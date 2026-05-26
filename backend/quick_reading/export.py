from __future__ import annotations

import html
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import quote

from backend.metadata import load_metadata
from backend.utils import safe_rel

from .models import (
    _Result,
    MISSING_TEXT,
    list_value,
    normalize_experiment_read,
    normalize_figure_insights,
    normalize_page_refs,
    normalize_technical_detail_read,
    text_value,
)
from .storage import load_speedread_cache, paper_cache_dir, paper_hash

ACCENT = "4f46e5"
TEXT = "1f2937"
MUTED = "64748b"
BORDER = "dbe3e1"
SURFACE = "ffffff"
SURFACE_SOFT = "f6f7fb"
FONT_STACK = '"Microsoft YaHei", "Segoe UI", Arial, sans-serif'


@dataclass
class ExportCard:
    title: str
    lines: list[tuple[str, str]] = field(default_factory=list)
    paragraphs: list[str] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    tags: list[str] = field(default_factory=list)
    page_refs: list[int] = field(default_factory=list)
    image_path: str = ""
    image_label: str = ""


@dataclass
class ExportSection:
    title: str
    paragraphs: list[str] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    cards: list[ExportCard] = field(default_factory=list)


@dataclass
class ExportDocument:
    file_key: str
    title: str
    metadata: list[str]
    generated_at: str
    model: str
    request_mode: str
    sections: list[ExportSection]


def _clean_text(value, fallback=MISSING_TEXT, limit=None):
    return text_value(value, fallback, limit)


def _clean_list(value, max_items=12):
    return list_value(value, max_items=max_items)


def _refs(value):
    return normalize_page_refs(value)


def _refs_text(refs):
    refs = _refs(refs)
    return f"Page {', '.join(str(item) for item in refs)}" if refs else ""


def _safe_image_path(image_path, workspace_root):
    if not image_path:
        return ""
    abs_path = safe_rel(str(image_path), workspace_root)
    if not abs_path or not os.path.isfile(abs_path):
        return ""
    return abs_path


def _paper_value(paper, key):
    value = paper.get(key) if isinstance(paper, dict) else ""
    return _clean_text(value, "", 600)


def _filename_piece(value, fallback):
    piece = re.sub(r"[\\/:*?\"<>|\r\n\t]+", "_", str(value or "")).strip(" ._")
    piece = re.sub(r"\s+", "_", piece)
    return (piece[:90] or fallback).strip("._")


def _download_filename(document, ext):
    return f"{_filename_piece(document.title, paper_hash(document.file_key))}_speedread.{ext}"


def _build_summary_section(content):
    return ExportSection(
        title="快速省流总结",
        paragraphs=[_clean_text(content.get("one_sentence"))],
        bullets=_clean_list(content.get("quick_takeaways"), 8),
    )


def _build_method_section(content):
    return ExportSection(
        title="方法速读",
        paragraphs=[_clean_text(content.get("method_overview"))],
        bullets=_clean_list(content.get("method_steps"), 10),
    )


def _build_technical_section(content):
    technical = normalize_technical_detail_read(content.get("technical_detail_read"))
    cards = []
    for item in technical.get("modules") or []:
        cards.append(ExportCard(
            title=_clean_text(item.get("name"), "技术模块"),
            lines=[
                ("作用", _clean_text(item.get("role"))),
                ("机制", _clean_text(item.get("mechanism"))),
                ("技术解释", _clean_text(item.get("technical_explanation"))),
            ],
            tags=_clean_list(item.get("terms"), 6),
            page_refs=_refs(item.get("page_refs")),
        ))
    for item in technical.get("terms") or []:
        basis = "通用背景" if item.get("basis") == "general_background" else "论文证据"
        cards.append(ExportCard(
            title=_clean_text(item.get("term"), "术语"),
            lines=[
                ("解释", _clean_text(item.get("explanation"))),
                ("本文上下文", _clean_text(item.get("paper_context"))),
                ("依据", basis),
            ],
        ))
    if technical.get("background_notes"):
        cards.append(ExportCard(title="必要背景", bullets=_clean_list(technical.get("background_notes"), 8)))
    if technical.get("missing_details"):
        cards.append(ExportCard(title="缺失细节", bullets=_clean_list(technical.get("missing_details"), 8)))
    return ExportSection(
        title="技术细节速读",
        paragraphs=[_clean_text(technical.get("summary"))],
        cards=cards,
    )


def _build_experiment_section(content):
    experiment = normalize_experiment_read(content.get("experiment_read"))
    cards = [ExportCard(
        title="复现概览",
        lines=[
            ("复现难度", {"low": "低", "medium": "中", "high": "高", "unknown": "未知"}.get(experiment.get("reproduction_difficulty"), "未知")),
            ("单人工作量", _clean_text(experiment.get("workload_estimate"))),
        ],
    )]
    for item in experiment.get("compared_methods") or []:
        cards.append(ExportCard(
            title=_clean_text(item.get("name"), "对比方法"),
            lines=[
                ("角色", _clean_text(item.get("role"))),
                ("说明", _clean_text(item.get("description"))),
                ("结果", _clean_text(item.get("result"))),
            ],
            page_refs=_refs(item.get("page_refs")),
        ))
    for item in experiment.get("datasets") or []:
        cards.append(ExportCard(
            title=_clean_text(item.get("name"), "数据集"),
            lines=[
                ("任务", _clean_text(item.get("task"))),
                ("规模", _clean_text(item.get("scale"))),
                ("划分", _clean_text(item.get("split"))),
                ("备注", _clean_text(item.get("notes"))),
            ],
            page_refs=_refs(item.get("page_refs")),
        ))
    for item in experiment.get("metrics") or []:
        cards.append(ExportCard(
            title=_clean_text(item.get("name"), "指标"),
            lines=[
                ("含义", _clean_text(item.get("meaning"))),
                ("使用原因", _clean_text(item.get("why_used"))),
                ("局限", _clean_text(item.get("limitation"))),
            ],
            page_refs=_refs(item.get("page_refs")),
        ))
    for item in experiment.get("main_results") or []:
        cards.append(ExportCard(
            title="主要结果",
            lines=[
                ("结论", _clean_text(item.get("claim"))),
                ("证据", _clean_text(item.get("evidence"))),
            ],
            page_refs=_refs(item.get("page_refs")),
        ))
    cards.extend([
        ExportCard(title="复现注意", bullets=_clean_list(experiment.get("reproduction_notes"), 8)),
        ExportCard(title="实验设计充分性", bullets=_clean_list(experiment.get("evaluation_design_notes"), 8)),
        ExportCard(title="缺失细节", bullets=_clean_list(experiment.get("missing_details"), 8)),
    ])
    return ExportSection(
        title="实验速读",
        paragraphs=[_clean_text(experiment.get("summary"))],
        cards=cards,
    )


def _build_figure_section(content, speed_read, workspace_root):
    cards = []
    options = speed_read.get("options") if isinstance(speed_read.get("options"), dict) else {}
    figure_items = normalize_figure_insights(
        content.get("figure_insights"),
        speed_read.get("candidate_pages") or [],
        options.get("depth") or "standard",
    )
    for item in figure_items:
        if not isinstance(item, dict):
            continue
        image_path = _safe_image_path(item.get("image_path"), workspace_root)
        page = item.get("page") or "?"
        label = _clean_text(item.get("label"), "关键图表")
        cards.append(ExportCard(
            title=_clean_text(item.get("title") or label, "关键图表"),
            paragraphs=[_clean_text(item.get("summary"))],
            lines=[("页码", f"Page {page}")],
            page_refs=_refs(item.get("page_refs") or [page]),
            image_path=image_path,
            image_label=label,
        ))
    return ExportSection(
        title="图表解读",
        paragraphs=[] if cards else ["当前速读没有图表解读。"],
        cards=cards,
    )


def build_export_document(file_key, workspace_root, metadata_file, quick_reading_cache_dir):
    if not file_key:
        return _Result(400, {"ok": False, "error": "缺少 file_key"})
    try:
        metadata = load_metadata(metadata_file)
    except Exception as exc:
        return _Result(500, {"ok": False, "error": f"读取 metadata 失败: {exc}"})
    paper = metadata.get(file_key)
    if not isinstance(paper, dict):
        return _Result(404, {"ok": False, "error": f"找不到论文: {file_key}"})
    try:
        speed_read = load_speedread_cache(file_key, workspace_root, quick_reading_cache_dir)
    except Exception as exc:
        return _Result(500, {"ok": False, "error": f"读取速读缓存失败: {exc}"})
    if not speed_read or speed_read.get("status") != "success" or not isinstance(speed_read.get("content"), dict):
        return _Result(404, {"ok": False, "error": "这篇论文还没有成功生成的速读缓存"})

    content = speed_read.get("content") or {}
    paper_snapshot = speed_read.get("paper") if isinstance(speed_read.get("paper"), dict) else {}
    merged_paper = {**paper, **paper_snapshot}
    metadata_bits = []
    venue_year = " ".join(item for item in [_paper_value(merged_paper, "venue"), _paper_value(merged_paper, "year")] if item)
    if venue_year:
        metadata_bits.append(venue_year)
    authors = _paper_value(merged_paper, "authors")
    if authors:
        metadata_bits.append(f"作者: {authors}")
    if speed_read.get("updated_at") or speed_read.get("generated_at"):
        metadata_bits.append(f"生成: {speed_read.get('updated_at') or speed_read.get('generated_at')}")
    if speed_read.get("model"):
        metadata_bits.append(f"模型: {speed_read.get('model')}")
    if speed_read.get("request_mode"):
        metadata_bits.append(f"模式: {speed_read.get('request_mode')}")

    sections = [
        _build_summary_section(content),
        ExportSection("问题与动机", paragraphs=[_clean_text(content.get("problem_and_motivation"))]),
        _build_method_section(content),
        _build_technical_section(content),
        _build_experiment_section(content),
        _build_figure_section(content, speed_read, workspace_root),
        ExportSection("贡献", bullets=_clean_list(content.get("contributions"), 8)),
        ExportSection("局限", bullets=_clean_list(content.get("limitations"), 8)),
        ExportSection("精读建议", bullets=_clean_list(content.get("deep_read_suggestions"), 10)),
    ]
    document = ExportDocument(
        file_key=file_key,
        title=_paper_value(merged_paper, "title") or file_key,
        metadata=metadata_bits,
        generated_at=str(speed_read.get("updated_at") or speed_read.get("generated_at") or ""),
        model=str(speed_read.get("model") or ""),
        request_mode=str(speed_read.get("request_mode") or ""),
        sections=sections,
    )
    return _Result(200, {"ok": True, "document": document})


def _html_text(value):
    return html.escape(str(value or MISSING_TEXT)).replace("\n", "<br>")


def _render_bullets(items):
    if not items:
        return ""
    return "<ul>" + "".join(f"<li>{_html_text(item)}</li>" for item in items) + "</ul>"


def _render_card(card):
    refs = _refs_text(card.page_refs)
    ref_html = f'<span class="page-ref">{_html_text(refs)}</span>' if refs else ""
    image_html = ""
    if card.image_path:
        image_html = f'<img class="figure-image" src="{Path(card.image_path).as_uri()}" alt="{_html_text(card.image_label or card.title)}">'
    elif card.image_label:
        image_html = '<div class="image-placeholder">截图缺失</div>'
    line_html = "".join(
        f'<p><b>{_html_text(label)}:</b> {_html_text(text)}</p>'
        for label, text in card.lines
        if text
    )
    paragraphs = "".join(f'<p>{_html_text(item)}</p>' for item in card.paragraphs if item)
    tags = "".join(f'<span class="tag">{_html_text(item)}</span>' for item in card.tags if item and item != MISSING_TEXT)
    tags_html = f'<div class="tags">{tags}</div>' if tags else ""
    return (
        '<article class="card">'
        f'{image_html}'
        f'<div class="card-head"><h3>{_html_text(card.title)}</h3>{ref_html}</div>'
        f'{line_html}{paragraphs}{_render_bullets(card.bullets)}{tags_html}'
        '</article>'
    )


def _render_print_html(document):
    meta = " · ".join(document.metadata)
    sections = []
    for section in document.sections:
        paragraphs = "".join(f'<p class="lead">{_html_text(item)}</p>' for item in section.paragraphs if item)
        cards = "".join(_render_card(card) for card in section.cards)
        sections.append(
            f'<section class="section"><h2>{_html_text(section.title)}</h2>{paragraphs}{_render_bullets(section.bullets)}'
            f'<div class="cards">{cards}</div></section>'
        )
    return f'''<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{_html_text(document.title)} - 速读导出</title>
<style>
@page {{ size: A4; margin: 16mm 14mm 18mm; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; font-family: {FONT_STACK}; color: #{TEXT}; background: #fff; font-size: 11pt; line-height: 1.62; }}
.cover {{ min-height: 210mm; display: flex; flex-direction: column; justify-content: center; border-bottom: 3px solid #{ACCENT}; page-break-after: always; }}
.cover-label {{ color: #{ACCENT}; font-weight: 700; letter-spacing: 0; font-size: 13pt; margin-bottom: 14px; }}
h1 {{ color: #111827; font-size: 27pt; line-height: 1.22; margin: 0 0 16px; letter-spacing: 0; }}
.meta {{ color: #{MUTED}; font-size: 10pt; line-height: 1.7; }}
.section {{ page-break-before: auto; margin: 0 0 14mm; }}
h2 {{ color: #111827; font-size: 18pt; line-height: 1.3; margin: 0 0 9mm; break-after: avoid; border-bottom: 1px solid #{BORDER}; padding-bottom: 4mm; }}
h3 {{ color: #111827; font-size: 12pt; margin: 0; line-height: 1.35; break-after: avoid; }}
p {{ margin: 0 0 6px; }}
.lead {{ color: #111827; font-size: 11.5pt; line-height: 1.72; margin-bottom: 8px; }}
ul {{ margin: 0 0 8px 0; padding-left: 18px; }}
li {{ margin: 0 0 5px; }}
.cards {{ display: block; }}
.card {{ border: 1px solid #{BORDER}; border-radius: 8px; background: #{SURFACE_SOFT}; padding: 10px 12px; margin: 0 0 9px; break-inside: avoid; page-break-inside: avoid; }}
.card-head {{ display: flex; justify-content: space-between; gap: 8px; align-items: flex-start; margin-bottom: 5px; }}
.card p {{ margin: 4px 0; }}
.page-ref, .tag {{ display: inline-flex; color: #{ACCENT}; background: #eef2ff; border-radius: 999px; padding: 1px 7px; font-size: 8.5pt; font-weight: 600; white-space: nowrap; }}
.tags {{ display: flex; flex-wrap: wrap; gap: 5px; margin-top: 6px; }}
.figure-image {{ width: 100%; max-height: 118mm; object-fit: contain; display: block; background: #eef2f7; border-radius: 6px; margin-bottom: 8px; }}
.image-placeholder {{ height: 62mm; display: flex; align-items: center; justify-content: center; color: #{MUTED}; background: #eef2f7; border-radius: 6px; margin-bottom: 8px; }}
</style>
</head>
<body>
<header class="cover">
<div class="cover-label">OpenPaper AI 速读导出</div>
<h1>{_html_text(document.title)}</h1>
<div class="meta">{_html_text(meta or 'metadata unavailable')}</div>
</header>
{''.join(sections)}
</body>
</html>'''


def _registry_app_path(names):
    if os.name != "nt":
        return ""
    try:
        import winreg
    except Exception:
        return ""
    roots = [winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER]
    subkeys = [fr"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\{name}" for name in names]
    for root in roots:
        for subkey in subkeys:
            try:
                with winreg.OpenKey(root, subkey) as key:
                    value, _ = winreg.QueryValueEx(key, "")
            except OSError:
                continue
            if value and os.path.isfile(value):
                return value
    return ""


def find_browser_executable():
    env_path = os.environ.get("OPENPAPER_BROWSER_PATH") or os.environ.get("EDGE_PATH") or os.environ.get("CHROME_PATH")
    if env_path and os.path.isfile(env_path):
        return env_path
    for name in ("msedge", "chrome", "chromium", "chromium-browser"):
        found = shutil.which(name)
        if found:
            return found
    candidates = []
    for env_name in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA"):
        base = os.environ.get(env_name)
        if not base:
            continue
        candidates.extend([
            os.path.join(base, "Microsoft", "Edge", "Application", "msedge.exe"),
            os.path.join(base, "Google", "Chrome", "Application", "chrome.exe"),
            os.path.join(base, "Chromium", "Application", "chrome.exe"),
        ])
    candidates.extend([
        "/usr/bin/google-chrome",
        "/usr/bin/chromium",
        "/usr/bin/chromium-browser",
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    ])
    for path in candidates:
        if path and os.path.isfile(path):
            return path
    return _registry_app_path(("msedge.exe", "chrome.exe"))


def _try_add_pdf_footer(pdf_path, document):
    try:
        import fitz
    except Exception:
        return
    try:
        doc = fitz.open(pdf_path)
        total = doc.page_count
        for index, page in enumerate(doc, start=1):
            rect = page.rect
            footer = f"OpenPaper AI 速读 · {index}/{total}"
            page.insert_text((40, rect.height - 20), footer, fontsize=8, color=(0.39, 0.45, 0.55))
        tmp_path = pdf_path + ".footer.pdf"
        doc.save(tmp_path, garbage=4, deflate=True)
        doc.close()
        os.replace(tmp_path, pdf_path)
    except Exception:
        try:
            doc.close()
        except Exception:
            pass


def _export_pdf(document, workspace_root, quick_reading_cache_dir):
    browser = find_browser_executable()
    if not browser:
        return _Result(500, {"ok": False, "error": "未找到 Edge/Chrome/Chromium，可安装浏览器或设置 OPENPAPER_BROWSER_PATH 后重试"})
    cache_dir = os.path.abspath(paper_cache_dir(document.file_key, workspace_root, quick_reading_cache_dir))
    os.makedirs(cache_dir, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="export_", dir=cache_dir) as temp_dir:
        html_path = os.path.abspath(os.path.join(temp_dir, "speed_read_export.html"))
        pdf_path = os.path.abspath(os.path.join(temp_dir, "speed_read_export.pdf"))
        with open(html_path, "w", encoding="utf-8") as file_obj:
            file_obj.write(_render_print_html(document))
        command = [
            browser,
            "--headless",
            "--disable-gpu",
            "--no-sandbox",
            "--disable-dev-shm-usage",
            "--allow-file-access-from-files",
            f"--print-to-pdf={pdf_path}",
            Path(html_path).as_uri(),
        ]
        try:
            completed = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
        except Exception as exc:
            return _Result(500, {"ok": False, "error": f"调用浏览器生成 PDF 失败: {exc}"})
        if completed.returncode != 0 or not os.path.isfile(pdf_path):
            detail = (completed.stderr or completed.stdout or "").strip()
            return _Result(500, {"ok": False, "error": f"浏览器生成 PDF 失败: {detail[:500] or '未输出错误详情'}"})
        _try_add_pdf_footer(pdf_path, document)
        with open(pdf_path, "rb") as file_obj:
            data = file_obj.read()
    return _Result(200, {
        "ok": True,
        "data": data,
        "filename": _download_filename(document, "pdf"),
        "content_type": "application/pdf",
    })


def _rgb(hex_value):
    value = hex_value.strip("#")
    return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))


def _set_text(run, size=14, bold=False, color=TEXT):
    from pptx.util import Pt
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _ppt_rgb(color)


def _ppt_rgb(hex_value):
    from pptx.dml.color import RGBColor
    return RGBColor(*_rgb(hex_value))


def _add_textbox(slide, left, top, width, height, lines, size=14, color=TEXT, bold_first=False, bullet=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, line in enumerate(lines or [MISSING_TEXT]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = ""
        paragraph.level = 0
        paragraph.space_after = 0
        if bullet:
            paragraph.text = f"• {line}"
            for run in paragraph.runs:
                _set_text(run, size=size, color=color)
        else:
            run = paragraph.add_run()
            run.text = str(line)
            _set_text(run, size=size, bold=bold_first and index == 0, color=color)
    return box


def _add_slide_title(slide, title, continued=False):
    from pptx.util import Inches
    suffix = "（续）" if continued else ""
    _add_textbox(slide, Inches(0.55), Inches(0.34), Inches(12.25), Inches(0.52), [title + suffix], size=22, color="111827", bold_first=True)
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(0.96), Inches(12.25), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = _ppt_rgb(ACCENT)
    line.line.fill.background()


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _estimate_card_units(card):
    text_len = len(card.title) + sum(len(label) + len(text) for label, text in card.lines) + sum(len(item) for item in card.paragraphs) + sum(len(item) for item in card.bullets)
    return max(3, min(9, 2 + text_len // 70 + len(card.lines) + len(card.bullets)))


def _split_section_chunks(section, max_units=18):
    entries = []
    if section.paragraphs:
        text_len = sum(len(item) for item in section.paragraphs)
        entries.append(("paragraphs", section.paragraphs, max(3, text_len // 85 + len(section.paragraphs))))
    if section.bullets:
        text_len = sum(len(item) for item in section.bullets)
        entries.append(("bullets", section.bullets, max(3, text_len // 80 + len(section.bullets))))
    for card in section.cards:
        entries.append(("card", card, _estimate_card_units(card)))

    chunks = []
    current = []
    units = 0
    for entry in entries:
        entry_units = entry[2]
        if current and units + entry_units > max_units:
            chunks.append(current)
            current = []
            units = 0
        current.append(entry)
        units += entry_units
    if current:
        chunks.append(current)
    return chunks or [[("paragraphs", [MISSING_TEXT], 2)]]


def _card_lines(card):
    lines = []
    if card.page_refs:
        lines.append(_refs_text(card.page_refs))
    for label, text in card.lines:
        if text:
            lines.append(f"{label}: {text}")
    lines.extend(card.paragraphs)
    if card.tags:
        lines.append("术语: " + " / ".join(item for item in card.tags if item and item != MISSING_TEXT))
    lines.extend(f"• {item}" for item in card.bullets if item)
    return lines or [MISSING_TEXT]


def _add_card(slide, card, left, top, width, height):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _ppt_rgb(SURFACE_SOFT)
    shape.line.color.rgb = _ppt_rgb(BORDER)
    _add_textbox(slide, left + Inches(0.16), top + Inches(0.1), width - Inches(0.32), Inches(0.28), [card.title], size=14, color="111827", bold_first=True)
    _add_textbox(slide, left + Inches(0.16), top + Inches(0.43), width - Inches(0.32), height - Inches(0.52), _card_lines(card), size=12.5, color=TEXT)


def _add_section_slides(prs, section):
    from pptx.util import Inches
    chunks = _split_section_chunks(section)
    for chunk_index, chunk in enumerate(chunks):
        slide = _blank_slide(prs)
        _add_slide_title(slide, section.title, continued=chunk_index > 0)
        y = Inches(1.22)
        for kind, payload, units in chunk:
            if kind == "paragraphs":
                height = Inches(max(0.55, min(1.7, 0.35 + units * 0.18)))
                _add_textbox(slide, Inches(0.7), y, Inches(11.9), height, payload, size=15, color="111827")
                y += height + Inches(0.12)
            elif kind == "bullets":
                height = Inches(max(0.75, min(4.8, 0.28 + units * 0.24)))
                _add_textbox(slide, Inches(0.82), y, Inches(11.4), height, payload, size=14, bullet=True)
                y += height + Inches(0.12)
            else:
                height = Inches(max(0.88, min(2.05, 0.38 + units * 0.19)))
                _add_card(slide, payload, Inches(0.7), y, Inches(11.9), height)
                y += height + Inches(0.12)


def _fit_image_size(image_path, max_width, max_height):
    try:
        from PIL import Image
        with Image.open(image_path) as image:
            width, height = image.size
        ratio = min(max_width / max(width, 1), max_height / max(height, 1))
        return int(width * ratio), int(height * ratio)
    except Exception:
        return max_width, max_height


def _add_figure_slides(prs, section):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    if not section.cards:
        _add_section_slides(prs, section)
        return
    for index, card in enumerate(section.cards):
        slide = _blank_slide(prs)
        _add_slide_title(slide, section.title, continued=index > 0)
        _add_textbox(slide, Inches(0.7), Inches(1.18), Inches(12), Inches(0.35), [card.title], size=16, color="111827", bold_first=True)
        if card.image_path:
            max_w = Inches(5.7)
            max_h = Inches(4.45)
            pic_w, pic_h = _fit_image_size(card.image_path, max_w, max_h)
            slide.shapes.add_picture(card.image_path, Inches(0.7), Inches(1.72), width=pic_w, height=pic_h)
        else:
            placeholder = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.72), Inches(5.7), Inches(4.0))
            placeholder.fill.solid()
            placeholder.fill.fore_color.rgb = _ppt_rgb("eef2f7")
            placeholder.line.color.rgb = _ppt_rgb(BORDER)
            _add_textbox(slide, Inches(2.45), Inches(3.38), Inches(2.2), Inches(0.4), ["截图缺失"], size=14, color=MUTED)
        _add_textbox(slide, Inches(6.65), Inches(1.72), Inches(5.75), Inches(4.6), _card_lines(card), size=14)


def _export_pptx(document, workspace_root, quick_reading_cache_dir):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except Exception as exc:
        return _Result(500, {"ok": False, "error": f"缺少 python-pptx/Pillow 依赖，请先安装 requirements.txt: {exc}"})

    cache_dir = os.path.abspath(paper_cache_dir(document.file_key, workspace_root, quick_reading_cache_dir))
    os.makedirs(cache_dir, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="export_", dir=cache_dir) as temp_dir:
        pptx_path = os.path.abspath(os.path.join(temp_dir, "speed_read_export.pptx"))
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        cover = _blank_slide(prs)
        _add_textbox(cover, Inches(0.75), Inches(1.1), Inches(11.9), Inches(0.42), ["OpenPaper AI 速读导出"], size=16, color=ACCENT, bold_first=True)
        _add_textbox(cover, Inches(0.75), Inches(1.72), Inches(11.9), Inches(1.75), [document.title], size=28, color="111827", bold_first=True)
        _add_textbox(cover, Inches(0.78), Inches(3.72), Inches(11.7), Inches(1.2), document.metadata or ["metadata unavailable"], size=14, color=MUTED)
        bar = cover.shapes.add_shape(1, Inches(0.75), Inches(6.45), Inches(11.9), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _ppt_rgb(ACCENT)
        bar.line.fill.background()

        for section in document.sections:
            if section.title == "图表解读":
                _add_figure_slides(prs, section)
            else:
                _add_section_slides(prs, section)
        prs.save(pptx_path)
        with open(pptx_path, "rb") as file_obj:
            data = file_obj.read()
    return _Result(200, {
        "ok": True,
        "data": data,
        "filename": _download_filename(document, "pptx"),
        "content_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    })


def export_speedread(file_key, export_format, workspace_root, metadata_file, quick_reading_cache_dir):
    fmt = (export_format or "").strip().lower()
    if fmt not in {"pdf", "pptx"}:
        return _Result(400, {"ok": False, "error": "format 必须是 pdf 或 pptx"})
    result = build_export_document(file_key, workspace_root, metadata_file, quick_reading_cache_dir)
    if result.status != 200:
        return result
    document = result.payload.get("document")
    if fmt == "pdf":
        return _export_pdf(document, workspace_root, quick_reading_cache_dir)
    return _export_pptx(document, workspace_root, quick_reading_cache_dir)


__all__ = ["build_export_document", "export_speedread", "find_browser_executable"]
